"""Carga rfp.xlsx → rfp_items (R-001…R-174) y anota el golden del slice.

Golden: R-001/R-038 (PDF), R-037 (Java/CAG settlement window), R-057 (abstención TPS).
R-018 y cualquier placeholder se filtran siempre en load-db / ensure_loaded.
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from rfp_agent.config import FORBIDDEN_GOLDEN_IDS, GOLDEN_IDS, RFP_DB, RFP_XLSX, RUNS_DB

SCHEMA_RFP = """
CREATE TABLE IF NOT EXISTS rfp_items (
    id TEXT PRIMARY KEY,
    texto TEXT NOT NULL,
    excel_row INTEGER NOT NULL,
    in_golden INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS golden (
    id TEXT PRIMARY KEY,
    respondible INTEGER NOT NULL,
    sql_verificacion TEXT NOT NULL,
    notas TEXT,
    FOREIGN KEY (id) REFERENCES rfp_items(id)
);
CREATE TABLE IF NOT EXISTS facts (
    id TEXT PRIMARY KEY,
    req_id TEXT NOT NULL,
    statement TEXT NOT NULL,
    source_doc TEXT,
    locator TEXT,
    kb_status TEXT,
    FOREIGN KEY (req_id) REFERENCES rfp_items(id)
);
"""

SCHEMA_RUNS = """
CREATE TABLE IF NOT EXISTS evidence_log (
    requirement_id TEXT NOT NULL,
    chunk_key TEXT NOT NULL,
    cita TEXT NOT NULL,
    created_at TEXT NOT NULL DEFAULT (datetime('now')),
    UNIQUE(requirement_id, chunk_key, cita)
);
CREATE TABLE IF NOT EXISTS run_trace (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    requirement_id TEXT NOT NULL,
    status TEXT,
    answer TEXT,
    payload TEXT,
    created_at TEXT NOT NULL DEFAULT (datetime('now'))
);
"""

_GOLDEN_ROWS_RAW = [
    {
        "id": "R-001",
        "respondible": 1,
        "sql_verificacion": "SELECT statement, source_doc, locator FROM facts WHERE req_id = 'R-001'",
        "notas": "Frase del PDF funcional, no del xlsx.",
    },
    {
        "id": "R-037",
        "respondible": 1,
        "sql_verificacion": "SELECT statement, source_doc, locator FROM facts WHERE req_id = 'R-037'",
        "notas": (
            "Código Java/CAG: SettlementWindowUtil (core_rtgs timetable). "
            "Hit retrieve scope=code; kb_status=CODIGO (no VALIDADO automático)."
        ),
    },
    {
        "id": "R-038",
        "respondible": 1,
        "sql_verificacion": "SELECT statement, source_doc, locator FROM facts WHERE req_id = 'R-038'",
        "notas": "Alineado con D02 de Taller 2: 24/7 y 24/7/365.",
    },
    {
        "id": "R-057",
        "respondible": 0,
        "sql_verificacion": "SELECT statement, source_doc, locator FROM facts WHERE req_id = 'R-057'",
        "notas": "El PDF no da TPS ni un máximo numérico. Abstenerse.",
    },
]

_FACTS_ROWS_RAW = [
    {
        "id": "F-001",
        "req_id": "R-001",
        "statement": "Liquidación de fondos final e irrevocable en tiempo real",
        "source_doc": "descripcion_funcional.pdf",
        "locator": "p.11",
        "kb_status": "VALIDADO",
    },
    {
        "id": "F-037",
        "req_id": "R-037",
        # Literal del hit retrieve_knowledge(scope=code): SettlementWindowUtil.java
        "statement": "Util class for settlement window",
        "source_doc": "SettlementWindowUtil.java",
        "locator": (
            "CORE_RTGS/com.montran.rtgs.tp/business/src/com/montran/rtgs/timetable/impl/"
            "SettlementWindowUtil.java"
        ),
        "kb_status": "CODIGO",
    },
    {
        "id": "F-038a",
        "req_id": "R-038",
        "statement": "El sistema RTGS de Montran se ha mejorado con capacidades para funcionar 24/7",
        "source_doc": "descripcion_funcional.pdf",
        "locator": "p.13",
        "kb_status": "VALIDADO",
    },
    {
        "id": "F-038b",
        "req_id": "R-038",
        "statement": "El RTGS de Montran puede funcionar y procesar pagos de forma 24/7/365",
        "source_doc": "descripcion_funcional.pdf",
        "locator": "p.15",
        "kb_status": "VALIDADO",
    },
]


def _es_placeholder(statement: str) -> bool:
    s = (statement or "").strip().lower()
    if not s:
        return True
    return (
        "exact phrase" in s
        or "[[rellenar]]" in s
        or s.startswith("<")
        or "copied from the pdf" in s
        or "frase literal" in s and "…" in s
    )


def _golden_rows_canonicos() -> list[dict]:
    allowed = set(GOLDEN_IDS)
    out = []
    for g in _GOLDEN_ROWS_RAW:
        rid = g["id"]
        if rid in FORBIDDEN_GOLDEN_IDS or rid not in allowed:
            continue
        out.append(g)
    return out


def _facts_rows_canonicos() -> list[dict]:
    allowed = set(GOLDEN_IDS)
    out = []
    for f in _FACTS_ROWS_RAW:
        rid = f["req_id"]
        if rid in FORBIDDEN_GOLDEN_IDS or rid not in allowed:
            continue
        if _es_placeholder(f.get("statement") or ""):
            continue
        out.append(f)
    return out


# Exportados para tests / drafting; siempre filtrados.
GOLDEN_ROWS = _golden_rows_canonicos()
FACTS_ROWS = _facts_rows_canonicos()


def connect_rfp(path: Path | None = None) -> sqlite3.Connection:
    db = path or RFP_DB
    db.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def connect_runs(path: Path | None = None) -> sqlite3.Connection:
    db = path or RUNS_DB
    db.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db))
    conn.row_factory = sqlite3.Row
    return conn


def _leer_xlsx(xlsx: Path) -> list[tuple[str, int, str]]:
    wb = load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = []
    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        texto = (row[0] or "").strip() if row else ""
        if not texto:
            continue
        rid = f"R-{len(filas) + 1:03d}"
        filas.append((rid, excel_row, texto))
    return filas


def _purge_forbidden(conn: sqlite3.Connection) -> list[str]:
    """Borra R-018 y cualquier golden fuera de GOLDEN_IDS. Devuelve ids eliminados."""
    removed: list[str] = []
    allowed = set(GOLDEN_IDS)
    for rid in list(FORBIDDEN_GOLDEN_IDS):
        if conn.execute("SELECT 1 FROM golden WHERE id = ?", (rid,)).fetchone():
            removed.append(rid)
        conn.execute("DELETE FROM facts WHERE req_id = ?", (rid,))
        conn.execute("DELETE FROM golden WHERE id = ?", (rid,))
        conn.execute("UPDATE rfp_items SET in_golden = 0 WHERE id = ?", (rid,))
    extras = [
        r[0]
        for r in conn.execute("SELECT id FROM golden").fetchall()
        if r[0] not in allowed
    ]
    for rid in extras:
        if rid not in removed:
            removed.append(rid)
        conn.execute("DELETE FROM facts WHERE req_id = ?", (rid,))
        conn.execute("DELETE FROM golden WHERE id = ?", (rid,))
        conn.execute("UPDATE rfp_items SET in_golden = 0 WHERE id = ?", (rid,))
    # Placeholders en facts
    for row in conn.execute("SELECT id, req_id, statement FROM facts").fetchall():
        if _es_placeholder(row["statement"]) or row["req_id"] not in allowed:
            conn.execute("DELETE FROM facts WHERE id = ?", (row["id"],))
            if row["req_id"] not in removed and row["req_id"] not in allowed:
                removed.append(row["req_id"])
    return removed


def load_rfp(xlsx: Path | None = None, db_path: Path | None = None) -> dict:
    xlsx = xlsx or RFP_XLSX
    if not xlsx.exists():
        raise FileNotFoundError(f"No está el workbook RFP: {xlsx}")
    filas = _leer_xlsx(xlsx)
    golden_rows = _golden_rows_canonicos()
    facts_rows = _facts_rows_canonicos()
    # Refrescar exports por si alguien mutó los raw en runtime.
    global GOLDEN_ROWS, FACTS_ROWS
    GOLDEN_ROWS = golden_rows
    FACTS_ROWS = facts_rows

    conn = connect_rfp(db_path)
    stripped: list[str] = []
    try:
        conn.executescript(SCHEMA_RFP)
        conn.execute("DELETE FROM facts")
        conn.execute("DELETE FROM golden")
        conn.execute("DELETE FROM rfp_items")
        for rid, excel_row, texto in filas:
            in_golden = int(rid in GOLDEN_IDS and rid not in FORBIDDEN_GOLDEN_IDS)
            conn.execute(
                "INSERT INTO rfp_items (id, texto, excel_row, in_golden) VALUES (?, ?, ?, ?)",
                (rid, texto, excel_row, in_golden),
            )
        for g in golden_rows:
            conn.execute(
                "INSERT INTO golden (id, respondible, sql_verificacion, notas) VALUES (?, ?, ?, ?)",
                (g["id"], g["respondible"], g["sql_verificacion"], g["notas"]),
            )
        for f in facts_rows:
            conn.execute(
                "INSERT INTO facts (id, req_id, statement, source_doc, locator, kb_status) VALUES (?, ?, ?, ?, ?, ?)",
                (f["id"], f["req_id"], f["statement"], f["source_doc"], f["locator"], f["kb_status"]),
            )
        stripped = _purge_forbidden(conn)
        conn.commit()
        golden_ids = [r[0] for r in conn.execute("SELECT id FROM golden ORDER BY id").fetchall()]
    finally:
        conn.close()
    return {
        "rfp_items": len(filas),
        "golden": len(golden_ids),
        "golden_ids": golden_ids,
        "facts": len(facts_rows),
        "stripped": stripped,
        "db": str(db_path or RFP_DB),
    }


def ensure_loaded() -> None:
    need_reload = False
    conn = connect_rfp()
    try:
        n = conn.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='rfp_items'"
        ).fetchone()[0]
        if n == 0:
            need_reload = True
        else:
            count = conn.execute("SELECT COUNT(*) FROM rfp_items").fetchone()[0]
            stripped = _purge_forbidden(conn)
            if stripped:
                conn.commit()
            golden_ids = {r[0] for r in conn.execute("SELECT id FROM golden").fetchall()}
            if count != 174 or golden_ids != set(GOLDEN_IDS):
                need_reload = True
    except sqlite3.Error:
        need_reload = True
    finally:
        conn.close()

    if need_reload:
        load_rfp()

    conn = connect_runs()
    try:
        conn.executescript(SCHEMA_RUNS)
        conn.commit()
    finally:
        conn.close()


def golden_snapshot() -> dict:
    """Estado actual del golden en SQLite (tras ensure_loaded)."""
    ensure_loaded()
    conn = connect_rfp()
    try:
        ids = [r[0] for r in conn.execute("SELECT id FROM golden ORDER BY id").fetchall()]
        n_facts = conn.execute("SELECT COUNT(*) FROM facts").fetchone()[0]
        return {
            "golden_count": len(ids),
            "golden_ids": ids,
            "facts": n_facts,
            "expected": list(GOLDEN_IDS),
            "ok": ids == list(GOLDEN_IDS) or set(ids) == set(GOLDEN_IDS),
        }
    finally:
        conn.close()


def get_item(req_id: str) -> dict | None:
    ensure_loaded()
    conn = connect_rfp()
    try:
        row = conn.execute("SELECT * FROM rfp_items WHERE id = ?", (req_id,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def list_golden() -> list[dict]:
    ensure_loaded()
    conn = connect_rfp()
    try:
        rows = conn.execute(
            "SELECT g.*, i.texto FROM golden g JOIN rfp_items i ON i.id = g.id ORDER BY g.id"
        ).fetchall()
        # Defensa en profundidad: nunca devolver R-018 al evaluador.
        return [dict(r) for r in rows if r["id"] in GOLDEN_IDS and r["id"] not in FORBIDDEN_GOLDEN_IDS]
    finally:
        conn.close()
