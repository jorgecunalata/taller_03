"""Índice RAG de documentos (PDF/DOCX/MD) → Qdrant `documentos`.

Alineado con Taller 2: chunking ~512/102, `kb_status`, piso de caracteres.
No vende el zip de taller02; reimplementa lo necesario para `index-docs`.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from rfp_agent.config import (
    CHUNK_TOKENS,
    COLLECTION_DOCS,
    CORPUS,
    DOC_EMBED_BATCH,
    DOC_EXTS,
    OVERLAP_TOKENS,
    OUT,
    PISO_CARACTERES_DOC,
    QDRANT_URL,
)
from rfp_agent.embedder import EmbeddingsUnavailable, backend_embeddings, embed
from rfp_agent.store import recrear, upsert

CHUNKS_PATH = OUT / "doc_chunks.jsonl"
INGESTA_PATH = OUT / "ingesta_documentos.json"


def _tokens_aprox(texto: str) -> list[str]:
    """Tokens aproximados (palabras) cuando no hay tokenizer de transformers."""
    return re.findall(r"\S+", texto or "")


def _unir_tokens(tokens: list[str]) -> str:
    return " ".join(tokens).strip()


def estado_kb(texto: str, nombre: str) -> str:
    """VALIDADO solo para corpus real; plantillas .md / [[RELLENAR]] no se promueven."""
    nombre_l = (nombre or "").lower()
    blob = texto or ""
    if nombre_l.endswith(".md"):
        return "EJEMPLO_NO_VALIDADO"
    if "EJEMPLO_NO_VALIDADO" in blob or "[[RELLENAR]]" in blob:
        return "EJEMPLO_NO_VALIDADO"
    if "NO_VALIDADO" in blob.upper() and "EJEMPLO" in blob.upper():
        return "EJEMPLO_NO_VALIDADO"
    return "VALIDADO"


def iter_doc_paths(corpus: Path | None = None) -> tuple[list[Path], dict]:
    """Archivos en `corpus/` y `corpus/docs/`. Omite directorios ausentes."""
    base = corpus or CORPUS
    dirs_info: dict[str, dict] = {
        "corpus": {
            "path": str(base),
            "exists": base.is_dir(),
            "skipped": not base.is_dir(),
        },
        "corpus/docs": {
            "path": str(base / "docs"),
            "exists": (base / "docs").is_dir(),
            "skipped": not (base / "docs").is_dir(),
        },
    }
    paths: list[Path] = []
    vistos: set[Path] = set()
    if not base.is_dir():
        return [], dirs_info

    for p in sorted(base.iterdir(), key=lambda x: x.name.lower()):
        if p.is_file() and p.suffix.lower() in DOC_EXTS:
            rp = p.resolve()
            if rp not in vistos:
                vistos.add(rp)
                paths.append(p)

    docs_dir = base / "docs"
    if docs_dir.is_dir():
        for p in sorted(docs_dir.rglob("*"), key=lambda x: str(x).lower()):
            if not p.is_file() or p.suffix.lower() not in DOC_EXTS:
                continue
            if any(part.startswith(".") for part in p.parts):
                continue
            rp = p.resolve()
            if rp in vistos:
                continue
            vistos.add(rp)
            paths.append(p)

    return paths, dirs_info


def _leer_pdf(path: Path) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    paginas = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            texto = page.extract_text() or ""
        except Exception:
            texto = ""
        paginas.append({"page": i, "text": texto})
    return paginas


def _leer_docx(path: Path) -> list[dict]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "Falta python-docx para indexar .docx. "
            "Instala: pip install python-docx"
        ) from exc
    doc = Document(str(path))
    partes: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            partes.append(t)
    for table in doc.tables:
        for row in table.rows:
            celdas = [(c.text or "").strip() for c in row.cells]
            linea = " | ".join(c for c in celdas if c)
            if linea:
                partes.append(linea)
    texto = "\n".join(partes)
    return [{"page": None, "text": texto}]


def _leer_texto(path: Path) -> list[dict]:
    texto = path.read_text(encoding="utf-8", errors="replace")
    return [{"page": None, "text": texto}]


def leer_documento(path: Path) -> tuple[list[dict], dict]:
    """Devuelve páginas {page, text} y fila de reporte de ingesta."""
    nombre = path.name
    suf = path.suffix.lower()
    if not path.exists():
        return [], {
            "archivo": nombre,
            "nombre": nombre,
            "estado": "ausente",
            "caracteres_utiles": 0,
            "paginas": 0,
            "chunks": 0,
        }
    try:
        if suf == ".pdf":
            paginas = _leer_pdf(path)
        elif suf == ".docx":
            paginas = _leer_docx(path)
        else:
            paginas = _leer_texto(path)
    except Exception as exc:
        return [], {
            "archivo": nombre,
            "nombre": nombre,
            "estado": f"error:{type(exc).__name__}",
            "detalle": str(exc)[:240],
            "caracteres_utiles": 0,
            "paginas": 0,
            "chunks": 0,
        }

    utiles = sum(len((p.get("text") or "").strip()) for p in paginas)
    estado = "indexado" if utiles >= PISO_CARACTERES_DOC else "rechazado_sin_texto_util"
    return paginas, {
        "archivo": nombre,
        "nombre": nombre,
        "estado": estado,
        "caracteres_utiles": utiles,
        "paginas": len(paginas),
        "chunks": 0,
    }


def fragmentar(paginas: list[dict], doc_id: str, kb_status: str) -> list[dict]:
    """Ventanas ~CHUNK_TOKENS con solape OVERLAP_TOKENS (Taller 2)."""
    paso = max(1, CHUNK_TOKENS - OVERLAP_TOKENS)
    chunks: list[dict] = []
    for pag in paginas:
        page = pag.get("page")
        tokens = _tokens_aprox(pag.get("text") or "")
        if not tokens:
            continue
        if len(tokens) <= CHUNK_TOKENS:
            ventanas = [tokens]
        else:
            ventanas = []
            for inicio in range(0, len(tokens), paso):
                pieza = tokens[inicio : inicio + CHUNK_TOKENS]
                if len(pieza) < 40 and ventanas:
                    break
                ventanas.append(pieza)
                if inicio + CHUNK_TOKENS >= len(tokens):
                    break
        locator = f"p.{page}" if page is not None else doc_id
        for n, pieza in enumerate(ventanas):
            texto = _unir_tokens(pieza)
            if len(texto.strip()) < 20:
                continue
            chunks.append(
                {
                    "chunk_key": f"{doc_id}::p{page}::{n}",
                    "doc_id": doc_id,
                    "page": page,
                    "locator": locator if n == 0 else f"{locator}#{n}",
                    "text": texto,
                    "kb_status": kb_status,
                    "tokens": len(pieza),
                    "source_type": "documento",
                    "language": "es",
                }
            )
    return chunks


def construir_chunks(corpus: Path | None = None) -> tuple[list[dict], dict]:
    paths, dirs_info = iter_doc_paths(corpus)
    base = corpus or CORPUS
    reporte: list[dict] = []
    todos: list[dict] = []
    pages_total = 0

    for path in paths:
        paginas, fila = leer_documento(path)
        try:
            fila["path"] = str(path.relative_to(base))
        except ValueError:
            fila["path"] = path.name
        if fila["estado"] != "indexado":
            reporte.append(fila)
            continue
        kb = estado_kb("\n".join(p.get("text") or "" for p in paginas), path.name)
        piezas = fragmentar(paginas, path.name, kb)
        fila["chunks"] = len(piezas)
        fila["kb_status"] = kb
        fila["tokens_chunk_max"] = max((p["tokens"] for p in piezas), default=0)
        pages_total += int(fila.get("paginas") or 0)
        reporte.append(fila)
        todos.extend(piezas)

    summary = {
        "corpus": str(base),
        "corpus_exists": base.is_dir(),
        "dirs": dirs_info,
        "files_seen": len(paths),
        "files_indexed": sum(1 for f in reporte if f.get("estado") == "indexado"),
        "files_rejected": sum(1 for f in reporte if f.get("estado") != "indexado"),
        "pages": pages_total,
        "chunks": len(todos),
        "chunk_tokens": CHUNK_TOKENS,
        "overlap_tokens": OVERLAP_TOKENS,
        "archivos": reporte,
    }
    return todos, summary


def indexar_documentos(
    corpus: Path | None = None,
    qdrant_url: str | None = None,
    recreate: bool = True,
) -> dict:
    """Construye chunks de corpus/ (+ corpus/docs/) y los sube a COLLECTION_DOCS."""
    OUT.mkdir(parents=True, exist_ok=True)
    chunks, summary = construir_chunks(corpus=corpus)

    if summary["files_seen"] == 0:
        summary["warning"] = (
            f"Ningún PDF/DOCX/MD bajo {summary['corpus']} ni corpus/docs/. "
            "Copia documentos a corpus/ (p. ej. descripcion_funcional.pdf)."
        )
        INGESTA_PATH.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return summary

    if not chunks:
        summary["warning"] = (
            "Se vieron archivos pero ninguno aportó texto útil "
            f"(piso {PISO_CARACTERES_DOC} caracteres)."
        )
        INGESTA_PATH.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return summary

    with CHUNKS_PATH.open("w", encoding="utf-8") as fh:
        for c in chunks:
            fh.write(json.dumps(c, ensure_ascii=False) + "\n")

    url = qdrant_url or QDRANT_URL
    try:
        if recreate:
            recrear(COLLECTION_DOCS, url=url)
        lote = DOC_EMBED_BATCH
        for i in range(0, len(chunks), lote):
            grupo = chunks[i : i + lote]
            vectores = embed([g["text"] for g in grupo], batch_size=lote)
            payloads = [{k: v for k, v in g.items() if k != "chunk_key"} for g in grupo]
            upsert(
                COLLECTION_DOCS,
                [g["chunk_key"] for g in grupo],
                vectores,
                payloads,
                lote=64,
                url=url,
            )
            if i % max(lote * 4, 1) == 0 or i + lote >= len(chunks):
                print(
                    f"  indexados {min(i + lote, len(chunks))}/{len(chunks)}",
                    flush=True,
                )
    except EmbeddingsUnavailable:
        raise
    except Exception as exc:
        raise RuntimeError(
            f"Fallo al escribir en Qdrant ({url} / {COLLECTION_DOCS}): {exc}\n"
            "En el Mac: docker run -p 6333:6333 -p 6334:6334 qdrant/qdrant"
        ) from exc

    summary.update(
        {
            "collection": COLLECTION_DOCS,
            "qdrant_url": url,
            "points_upserted": len(chunks),
            "embedding_backend": backend_embeddings(),
            "chunks_path": str(CHUNKS_PATH),
            "ingesta_path": str(INGESTA_PATH),
        }
    )
    INGESTA_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def imprimir_resumen(summary: dict) -> None:
    print(f"CORPUS={summary.get('corpus')} exists={summary.get('corpus_exists')}")
    print(
        f"colección={summary.get('collection', COLLECTION_DOCS)} "
        f"url={summary.get('qdrant_url', QDRANT_URL)}"
    )
    dirs = summary.get("dirs") or {}
    for key in ("corpus", "corpus/docs"):
        d = dirs.get(key) or {}
        if d.get("skipped") or not d.get("exists"):
            print(f"  {key}: SKIP (missing)")
        else:
            print(f"  {key}: ok path={d.get('path')}")
    for fila in summary.get("archivos") or []:
        estado = fila.get("estado")
        print(
            f"  {fila.get('path') or fila.get('nombre')}: "
            f"estado={estado} pages={fila.get('paginas', 0)} "
            f"chunks={fila.get('chunks', 0)} kb={fila.get('kb_status', '-')}"
        )
    print(
        f"total files={summary.get('files_indexed', 0)}/{summary.get('files_seen', 0)} "
        f"pages={summary.get('pages', 0)} "
        f"chunks={summary.get('chunks', 0)} "
        f"points={summary.get('points_upserted', 0)} "
        f"embed={summary.get('embedding_backend', '?')}"
    )
    if summary.get("warning"):
        print(f"WARNING: {summary['warning']}")
